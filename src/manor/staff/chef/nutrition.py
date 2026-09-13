"""食品成分表から材料の栄養値を推定する（ADR-019）。

**数字は表から引き、LLM は呼ばない**——文部科学省「日本食品標準成分表（八訂）増補
2023 年」の写し（`chef_food`）と、材料名 → 食品番号の名寄せ（`chef_food_alias`）と、
`lexicon.toml` の `[units]`（分量 → グラム）だけで組み立てる。`claude -p` の出番は
第三段（名寄せの下書き）で、この module には入っていない。

## 立て付け（`menu.py` と同じ）

- 正規化・名寄せ・換算・合算 … **純粋関数**（DB も外部も触らない。合成 dict で試験できる）
- `import_food_table()` / `refresh()` / `rebuild()` … DB を読み書きする薄い層

CLI（`cli.py`）・Web（`web/api_v1/kitchen.py`）の両方がここの関数を呼ぶ。数値は1つも
持たない——換算の物差しは全部 `lexicon.toml` の `[units]`・`[food_normalize]` から読む
（`menu.py` が帯・重みを持たないのと同じ約束。主人が書き換えれば結果が変わる）。

## 名寄せの4段（ADR-019 D2）

1. `chef_food_alias.alias` の**完全一致**（素の名前 → 正規化した名前の順に引く）
2. 正規化した名前が**成分表の食品名と完全一致**
3. 成分表の食品名に**部分一致**（当たり方の**位置**で順位を付け、同じ位置なら「生」を
   優先し、次に調理済みでないもの、最後に名前が短いもの。`_match_tier` 参照）
4. 当たらなければ**未解決**——推定値には足さず、`coverage` を下げる

①の alias は主人が手で入れるほか、**名寄せの種**（`food_aliases_seed.toml`。ADR-019 §4）
が `confidence='rule'` で先に埋める——成分表の食品名は「根深ねぎ 葉 軟白 生」のような
分類の言葉で、うちの「長ねぎ」とは字が重ならない。字が無いものは③では当たらない。

## coverage（解決率）と `partial`

`coverage` は**重量比**（解決できた材料のグラム数 ÷ 数えた材料のグラム数）。換算が
できなかった材料は `[units].unresolved_grams`（目安）として**分母にだけ**数える
——分からないものを 0 と見なすと coverage が甘くなる。`少々`・`適量` で食品ごとの
固定値を持たないものは**どちらにも数えない**（ADR-019 D3「数えない」）。

`coverage` が `[menu.rules].nutrition_coverage_min`（既定 0.8）を下回るものは `partial`
で、献立の候補には入らない（ADR-018 D1 に足した規則。`menu.py` が見る）。

## 上書きしない約束（ADR-019 D4）

`nutrition_source` が `site`（出典サイトの表示値）・`manual`（手入力）のレシピは
**触らない**。推定が書くのは空欄と、以前の推定（`estimated`）だけ。
"""

from __future__ import annotations

import csv
import sqlite3
import re
import tomllib
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from manor import util
from manor.errors import ManorError

from . import ops, recipe_shaping as shaping, recipes

#: 栄養の5項目（`chef_recipe_meta`・`chef_food` の列名。ADR-019 D1）。
NUTRIENTS: tuple[str, ...] = ("kcal", "protein_g", "fat_g", "carb_g", "salt_g")

#: 成分表の版（`chef_food.source_version` の既定）。
DEFAULT_SOURCE_VERSION = "8th-2023"

#: `chef_food.per`。100g あたりの値だけを持つ（ADR-019 D1）。
PER_100G = "100g"

#: `chef_food_alias.confidence` の語彙（`schema.sql` の CHECK と一致させる）。
VALID_CONFIDENCE: tuple[str, ...] = ("manual", "rule", "llm")

#: 名寄せがどの段で当たったか（ADR-019 D2）。`""` は未解決。
STAGE_ALIAS, STAGE_NAME, STAGE_PARTIAL = "alias", "name", "partial"

#: 未解決の理由（画面が文へ直す符牒。`nutrition.py` は文を組まない——`menu.py` の
#: 「理由は定型文の符牒で返す」と同じ判断）。
REASON_NO_FOOD = "no_food"        # 分量は分かったが成分表に当たらない
REASON_NO_AMOUNT = "no_amount"    # 分量が書かれていない
REASON_UNKNOWN_UNIT = "unknown_unit"  # 単位を換算できない
REASON_NO_PIECE = "no_piece"      # 「個」「本」だが、その食品の目安重量を持たない

#: `to_grams()` が「数えない」を返すときの印（ADR-019 D3。`適量`・対象外の `少々`）。
NOT_COUNTED = "not_counted"

#: 取り込み・名寄せの誤りのキー（web が状態コードへ写す。`media.py` と同じ流儀）。
ERR_FOOD_TABLE_MISSING = "error.chef.food_table_missing"
ERR_FOOD_IMPORT_FAILED = "error.chef.food_import_failed"
ERR_FOOD_NOT_FOUND = "error.chef.food_not_found"

#: 部分一致で「生」を優先するための語（ADR-019 D2 ③）。
_RAW_WORD = "生"
#: 調理済みを示す語（部分一致の優先度を下げる）。
_COOKED_WORDS: tuple[str, ...] = (
    "ゆで", "茹で", "焼き", "蒸し", "油いため", "水煮", "乾", "フライ", "から揚げ", "天ぷら", "缶詰",
)

#: 名寄せの種（ADR-019 §4）。`lexicon.toml` と同じ「語彙はファイルに置く」流儀。
_SEED_PATH = Path(__file__).with_name("food_aliases_seed.toml")


# --- lexicon（換算の物差しの唯一の出どころ。ここは読むだけ） -----------------------------


@dataclass(frozen=True)
class UnitTables:
    """`lexicon.toml` の `[units]`・`[food_normalize]` の写し（ADR-019 D3）。"""

    volume_ml: dict[str, float] = field(default_factory=dict)
    weight_g: dict[str, float] = field(default_factory=dict)
    density: dict[str, float] = field(default_factory=dict)
    piece: dict[str, dict[str, float]] = field(default_factory=dict)
    pinch: dict[str, dict[str, float]] = field(default_factory=dict)
    unresolved_grams: float = 30.0
    drop_words: tuple[str, ...] = ()
    synonyms: tuple[tuple[str, str], ...] = ()


def _float_map(raw: Any) -> dict[str, float]:
    out: dict[str, float] = {}
    for key, value in dict(raw or {}).items():
        try:
            out[str(key)] = float(value)
        except (TypeError, ValueError):
            continue
    return out


def load_unit_tables(path: Path | None = None) -> UnitTables:
    """`[units]`・`[food_normalize]` を読む（試験は `path` に合成データを渡せる）。"""
    lex = ops.load_lexicon(path)
    units = dict(lex.get("units") or {})  # type: ignore[arg-type]
    normalize = dict(lex.get("food_normalize") or {})  # type: ignore[arg-type]
    piece_raw = dict(units.get("piece") or {})
    pinch_raw = dict(units.get("pinch") or {})
    try:
        unresolved = float(units.get("unresolved_grams", 30.0))
    except (TypeError, ValueError):
        unresolved = 30.0
    # 同義語は**長い鍵を先に**当てる（「豚こま切れ肉」を「豚こま」より先に見る）。
    synonyms = tuple(
        sorted(
            ((str(k), str(v)) for k, v in dict(normalize.get("synonyms") or {}).items()),
            key=lambda kv: -len(kv[0]),
        )
    )
    drop_words = tuple(sorted((str(w) for w in (normalize.get("drop_words") or [])), key=lambda w: -len(w)))
    return UnitTables(
        volume_ml=_float_map(units.get("volume_ml")),
        weight_g=_float_map(units.get("weight_g")),
        density=_float_map(units.get("density")),
        piece={str(k): _float_map(v) for k, v in piece_raw.items()},
        pinch={str(k): _float_map(v) for k, v in pinch_raw.items()},
        unresolved_grams=unresolved,
        drop_words=drop_words,
        synonyms=synonyms,
    )


def coverage_min(path: Path | None = None) -> float:
    """`partial` の境目（`[menu.rules].nutrition_coverage_min`）。

    献立の候補を絞る規則（ADR-018 D1）と同じ数字なので、**同じ節から読む**
    ——`menu.py` は `load_rules()` 経由で、ここは直に。写しを2つ持たない。
    """
    lex = ops.load_lexicon(path)
    rules = dict((lex.get("menu") or {}).get("rules") or {})  # type: ignore[union-attr]
    try:
        return float(rules.get("nutrition_coverage_min", 0.8))
    except (TypeError, ValueError):
        return 0.8


# --- 正規化（純粋関数。ADR-019 D2 ②） ----------------------------------------------------


def normalize_name(name: str, tables: UnitTables | None = None, *, drop: bool = True) -> str:
    """材料名 → 正規化した名前（`chef_food_alias.alias` に入れる形）。

    `recipe_shaping.normalize_food_name()`（機械的な均し）を通してから、
    `lexicon.toml` の `[food_normalize]` の語彙で下ごしらえ語を落とし、同義語へ寄せる。
    削り切って空になったら、削る前の形へ戻す（名前を失うより残すほうがよい）。

    `drop=False` は**下ごしらえ語を落とさない**——成分表の食品名（`build_index`）に
    使う。成分表の名前は「こまつな 葉 ゆで」のように**調理法が食品の区別そのもの**
    なので、材料名と同じ削りを掛けると「生」と「ゆで」が同じ名前へ潰れてしまう
    （どちらを採るかは `_partial_sort_key` が「生を優先」で決める。ADR-019 D2 ③）。
    同義語（かなの揺れ）は**両側に掛ける**——片側だけでは寄せた意味がない。
    """
    tables = tables or load_unit_tables()
    base = shaping.normalize_food_name(name)
    if not base:
        return ""
    text = base
    if drop:
        for word in tables.drop_words:
            if word:
                text = text.replace(shaping.normalize_food_name(word), "")
        text = text.strip() or base
    for src, dst in tables.synonyms:
        key = shaping.normalize_food_name(src)
        target = shaping.normalize_food_name(dst)
        if not key or key not in text:
            continue
        if target and target in text:
            # 既に寄せ先の形（「長ねぎ」に「ねぎ」→「長ねぎ」を当てて「長長ねぎ」に
            # するような二重適用）を防ぐ。同義語は**1 方向・1 回だけ**。
            continue
        text = text.replace(key, target)
    return text.strip() or base


# --- 分量 → グラム（純粋関数。ADR-019 D3） -----------------------------------------------

#: 数字（全角・分数・帯分数・範囲）。`recipe_shaping._NUM_PART` と同じ語彙を
#: **数値として読む**ためにここで解く（あちらは切り出すだけ）。
_RANGE_CHARS = ("〜", "~", "ー", "-", "–")


def parse_number(text: str) -> float | None:
    """`"2"` `"1/2"` `"2と1/2"` `"２"` `"2〜3"` を数値にする。読めなければ `None`。

    範囲（`2〜3`）は**平均**を採る——どちらかに寄せる理由が無く、平均のほうが
    推定の誤差が小さい。
    """
    raw = unicodedata.normalize("NFKC", str(text or "")).strip()
    if not raw:
        return None
    for ch in _RANGE_CHARS:
        if ch in raw[1:]:  # 先頭のマイナスは範囲ではない
            parts = [p for p in raw.split(ch) if p.strip()]
            values = [v for v in (parse_number(p) for p in parts) if v is not None]
            if values:
                return sum(values) / len(values)
            return None
    total = 0.0
    ok = False
    for part in raw.split("と"):
        part = part.strip()
        if not part:
            continue
        try:
            if "/" in part:
                num, _, den = part.partition("/")
                total += float(num) / float(den)
            else:
                total += float(part)
            ok = True
        except (TypeError, ValueError, ZeroDivisionError):
            return None
    return total if ok else None


def _lookup_by_name(table: Mapping[str, float], normalized: str) -> float | None:
    """食品名の表（比重・目安重量）から、名前に当たる値を1つ。

    突き合わせは**含む／含まれる**だけ（`ops.item_match` の「先頭と末尾の字が一致」は
    材料名向けに緩すぎて、「鮭」と「酒」のような別物まで当たる）。複数当たったら
    **鍵が長いもの**を採る（「ごま油」を「油」より優先する）。
    """
    best: tuple[int, float] | None = None
    for key, value in table.items():
        k = shaping.normalize_food_name(key)
        if not k:
            continue
        if k in normalized or normalized in k:
            if best is None or len(k) > best[0]:
                best = (len(k), value)
    return None if best is None else best[1]


def density_of(normalized: str, tables: UnitTables) -> float:
    """比重（g/ml）。表に無ければ 1.0（水と同じと見なす。ADR-019 D3）。"""
    value = _lookup_by_name(tables.density, normalized)
    return 1.0 if value is None else value


def to_grams(
    qty: str, unit: str, normalized: str, tables: UnitTables
) -> tuple[float | None, str]:
    """1つの材料の分量をグラムにする（ADR-019 D3）。

    戻り値 `(grams, reason)`:

    - `(数値, "")` … 換算できた
    - `(0.0, NOT_COUNTED)` … 「適量」等で**数えない**（分母にも入れない）
    - `(None, 理由)` … 換算できない（未解決。`coverage` の分母にだけ数える）

    `normalized` は `normalize_name()` を通した材料名（比重・目安重量を引くのに使う）。
    """
    unit = unicodedata.normalize("NFKC", str(unit or "")).strip()
    number = parse_number(qty)

    # ① 数字を伴わない量（少々・ひとつまみ）。食品ごとの固定値があればそれ、無ければ 0。
    if unit in tables.pinch:
        value = _lookup_by_name(tables.pinch[unit], normalized)
        return (value, "") if value is not None else (0.0, NOT_COUNTED)
    if not unit and number is None:
        return None, REASON_NO_AMOUNT
    if unit and number is None and unit not in tables.pinch:
        # 「適量」「お好みで」のように量を持たない語（`[units.pinch]` に節が無い）。
        if unit in tables.weight_g or unit in tables.volume_ml or unit in tables.piece:
            return None, REASON_NO_AMOUNT
        return 0.0, NOT_COUNTED
    assert number is not None  # noqa: S101 - 上の分岐で None は返し終えている
    if not unit:
        return None, REASON_UNKNOWN_UNIT

    # ② 重さ
    if unit in tables.weight_g:
        return number * tables.weight_g[unit], ""

    # ③ 容量（比重で補正）
    if unit in tables.volume_ml:
        return number * tables.volume_ml[unit] * density_of(normalized, tables), ""

    # ④ 個・本・枚・片・束…（食品ごとの目安重量）
    if unit in tables.piece:
        per = _lookup_by_name(tables.piece[unit], normalized)
        if per is None:
            return None, REASON_NO_PIECE
        return number * per, ""

    return None, REASON_UNKNOWN_UNIT


# --- 名寄せ（純粋関数。ADR-019 D2） ------------------------------------------------------


@dataclass
class FoodIndex:
    """成分表の引き方を1つにまとめた索引（純粋関数が DB を触らずに済むように）。"""

    by_code: dict[str, dict[str, Any]] = field(default_factory=dict)
    by_name: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    normalized: list[tuple[str, dict[str, Any]]] = field(default_factory=list)
    #: 食品番号 → 食品名を**語に割ったもの**（`_name_words`）。順位付けだけに使う。
    words: dict[str, tuple[str, ...]] = field(default_factory=dict)


#: 食品名の語の区切り（八訂は「ぶた ばら 脂身つき 生」のように**空白で語を並べる**）。
_NAME_WORD_SPLIT_RE = re.compile(r"[\s　]+")
#: 食品群の見出し（`<調味料類>`）。語ではなく分類の札なので、順位付けからは外す
#: （`(しょうゆ類)`・`[大型種肉]` は `recipe_shaping` の括弧落としが既に外している）。
_GROUP_LABEL_RE = re.compile(r"<[^>]*>")


def _name_words(name: str, tables: UnitTables) -> tuple[str, ...]:
    """成分表の食品名 → 正規化した**語の並び**（ADR-019 §4）。

    `normalize_name()` は空白を詰めて1本の文字列にしてしまうので、語の切れ目が消える
    ——「そらまめ しょうゆ豆」と「こいくちしょうゆ」のどちらが**しょうゆそのもの**かは、
    語の切れ目が分からないと決められない。ここだけ**割る前**の名前から語を採る。
    """
    raw = unicodedata.normalize("NFKC", str(name or ""))
    out: list[str] = []
    for token in _NAME_WORD_SPLIT_RE.split(raw):
        token = _GROUP_LABEL_RE.sub("", token).strip()
        if not token:
            continue
        word = normalize_name(token, tables, drop=False)
        if word:
            out.append(word)
    return tuple(out)


def build_index(foods: Iterable[Mapping[str, Any]], tables: UnitTables | None = None) -> FoodIndex:
    """`chef_food` の行（または合成 dict）から索引を組む。"""
    tables = tables or load_unit_tables()
    index = FoodIndex()
    for raw in foods:
        row = dict(raw)
        code = str(row.get("food_code") or "")
        if not code:
            continue
        index.by_code[code] = row
        norm = normalize_name(str(row.get("name") or ""), tables, drop=False)
        if not norm:
            continue
        index.by_name.setdefault(norm, []).append(row)
        index.normalized.append((norm, row))
        index.words[code] = _name_words(str(row.get("name") or ""), tables)
    return index


#: 部分一致の**当たり方**（ADR-019 §4。小さいほど「その食品そのもの」に近い）。
#: 日本語の複合語は**後ろが主辞**（「こいくち＋しょうゆ」はしょうゆの一種、
#: 「しょうゆ＋豆」は豆の一種）なので、語の**末尾**に当たったものを上に置く。
MATCH_WORD_EXACT = 0   # 食品名の語がまるごと材料名（「にんじん 根 皮つき 生」の「にんじん」）
MATCH_WORD_HEAD = 1    # 語の主辞が材料名（「こいくちしょうゆ」「食塩」「りょくとうもやし」）
MATCH_SPAN = 2         # 語の切れ目をまたいで含む（素の部分一致。「鶏 ひき肉 生」×「鶏ひき肉」）
MATCH_WORD_PREFIX = 3  # 語の**頭**にだけ当たる（「しょうゆ豆」「しょうゆ漬」＝別の食品の小分類）
MATCH_WORD_INSIDE = 4  # 語の**途中**にだけ当たる（「塩蔵わかめ」の「塩」）


def _match_tier(query: str, words: Sequence[str]) -> int:
    """材料名が食品名の**どこ**に当たったか（上の `MATCH_*`）。

    `query` が空（食品名のほうが材料名に含まれる向きの照合）なら、位置では区別せず
    `MATCH_SPAN` を返す——順位は「生」優先と名前の短さだけで決める（従来どおり）。
    """
    if not query or not words:
        return MATCH_SPAN
    tier: int | None = None
    for word in words:
        if word == query:
            return MATCH_WORD_EXACT
        if query not in word:
            continue
        if word.endswith(query):
            here = MATCH_WORD_HEAD
        elif word.startswith(query):
            here = MATCH_WORD_PREFIX
        else:
            here = MATCH_WORD_INSIDE
        tier = here if tier is None else min(tier, here)
    # どの語の中にも収まらない＝語の切れ目をまたいで含まれている。
    return MATCH_SPAN if tier is None else tier


def _partial_sort_key(
    norm_name: str,
    row: Mapping[str, Any],
    *,
    query: str = "",
    words: Sequence[str] = (),
) -> tuple[int, int, int, str]:
    """部分一致の優先度（ADR-019 D2 ③・§4）。

    ①当たり方の位置（`_match_tier`）→ ②「生」を含む → ③調理済みの語を含まない →
    ④それ以外、の順。同じ段なら**名前が短いもの**（余計な修飾が付いていない＝素の
    食品に近い）、最後に食品番号で安定させる。

    名前の長さは**食品群の見出しを除いた語の合計**で測る——`<調味料類>` の 6 字が
    調味料だけに一律で乗ると、群をまたいだ比べ（「そらまめ しょうゆ豆」対
    「こいくちしょうゆ」）が歪む。
    """
    if _RAW_WORD in norm_name:
        rank = 0
    elif not any(w in norm_name for w in _COOKED_WORDS):
        rank = 1
    else:
        rank = 2
    length = sum(len(w) for w in words) if words else len(norm_name)
    return _match_tier(query, words), rank, length, str(row.get("food_code") or "")


def resolve_food(
    name: str,
    index: FoodIndex,
    aliases: Mapping[str, str] | None = None,
    tables: UnitTables | None = None,
) -> dict[str, Any] | None:
    """材料名 → 成分表の1行（ADR-019 D2 の4段）。当たらなければ `None`。

    戻り値には `_stage`（どの段で当たったか）と `_normalized`（正規化した名前）を添える
    ——画面が「どうしてこの食品になったか」を見せられるようにする。
    """
    tables = tables or load_unit_tables()
    aliases = dict(aliases or {})
    norm = normalize_name(name, tables)
    if not norm:
        return None

    # ① alias の完全一致（素の名前でも引く——画面から手で入れた alias は素の形のことがある）
    for key in (shaping.normalize_food_name(name), norm):
        code = aliases.get(key)
        if code and code in index.by_code:
            return {**index.by_code[code], "_stage": STAGE_ALIAS, "_normalized": norm}

    # ② 正規化して成分表の食品名と完全一致
    exact = index.by_name.get(norm)
    if exact:
        row = sorted(exact, key=lambda r: _partial_sort_key(norm, r))[0]
        return {**row, "_stage": STAGE_NAME, "_normalized": norm}

    # ③ 部分一致（成分表の名前が材料名を含む → 材料名が成分表の名前を含む）
    # 前者は**当たり方の位置**で順位を付けられる（`_match_tier`）。後者は材料名のほうが
    # 長い＝食品名の語の話ではないので、位置では区別しない（`query=""` を渡す）。
    contains = [(n, r) for n, r in index.normalized if norm in n]
    query = norm
    if not contains:
        contains = [(n, r) for n, r in index.normalized if n in norm]
        query = ""
    if contains:
        n, row = sorted(
            contains,
            key=lambda pair: _partial_sort_key(
                pair[0],
                pair[1],
                query=query,
                words=index.words.get(str(pair[1].get("food_code") or ""), ()),
            ),
        )[0]
        return {**row, "_stage": STAGE_PARTIAL, "_normalized": norm}

    # ④ 未解決
    return None


# --- 推定（純粋関数。ADR-019 D4） --------------------------------------------------------


@dataclass
class Estimate:
    """1 レシピの推定（ADR-019 D4）。`nutrition` は**1 人前**。"""

    nutrition: dict[str, float] = field(default_factory=dict)
    coverage: float = 0.0
    servings: int = 1
    resolved: list[dict[str, Any]] = field(default_factory=list)
    unresolved: list[dict[str, Any]] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        """1つでも材料を解決できたか（0 なら `chef_recipe_meta` には書かない）。"""
        return bool(self.resolved)

    def is_partial(self, minimum: float | None = None) -> bool:
        """`coverage` が閾値を下回る（＝献立の候補に入れない）か。"""
        return self.coverage < (coverage_min() if minimum is None else minimum)

    def as_dict(self) -> dict[str, Any]:
        return {
            **{key: self.nutrition.get(key) for key in NUTRIENTS},
            "coverage": round(self.coverage, 3),
            "servings": self.servings,
            "resolved": self.resolved,
            "unresolved": self.unresolved,
        }


def estimate_nutrition(
    recipe: Mapping[str, Any],
    index: FoodIndex,
    aliases: Mapping[str, str] | None = None,
    tables: UnitTables | None = None,
) -> Estimate:
    """材料と分量から1人前の栄養値を推定する（ADR-019 D4）。**DB も外部も触らない。**

    材料ごとに `グラム × 成分/100` を足し、**廃棄率を掛け**（買った重量から可食部を
    出す）、`servings` で割って1人前にする。`servings` が無ければ 1 人前と見なす
    ——割らないほうが「1 皿ぶん」として意味が通る。

    `recipe` は `recipes.get()` の形（`servings` と `ingredients[]` を読む）。
    """
    tables = tables or load_unit_tables()
    aliases = dict(aliases or {})
    try:
        servings = max(1, int(recipe.get("servings") or 1))
    except (TypeError, ValueError):
        servings = 1

    totals = {key: 0.0 for key in NUTRIENTS}
    resolved_grams = 0.0
    counted_grams = 0.0
    resolved: list[dict[str, Any]] = []
    unresolved: list[dict[str, Any]] = []

    for ing in recipe.get("ingredients") or []:
        if not isinstance(ing, Mapping):
            continue
        name = str(ing.get("name") or "").strip()
        if not name:
            continue
        norm = normalize_name(name, tables)
        grams, reason = to_grams(str(ing.get("qty") or ""), str(ing.get("unit") or ""), norm, tables)
        food = resolve_food(name, index, aliases, tables)

        if grams is None:
            # 換算できない。分母にだけ目安重量で数える（分からないほど coverage が下がる）。
            counted_grams += tables.unresolved_grams
            unresolved.append(
                {
                    "name": name,
                    "normalized": norm,
                    "qty": str(ing.get("qty") or ""),
                    "unit": str(ing.get("unit") or ""),
                    "reason": reason,
                }
            )
            continue
        if reason == NOT_COUNTED:
            continue  # 「適量」等は**数えない**（ADR-019 D3）

        counted_grams += grams
        if food is None:
            unresolved.append(
                {
                    "name": name,
                    "normalized": norm,
                    "qty": str(ing.get("qty") or ""),
                    "unit": str(ing.get("unit") or ""),
                    "grams": round(grams, 1),
                    "reason": REASON_NO_FOOD,
                }
            )
            continue

        try:
            refuse = float(food.get("refuse_pct") or 0.0)
        except (TypeError, ValueError):
            refuse = 0.0
        edible = grams * max(0.0, 1.0 - refuse / 100.0)
        for key in NUTRIENTS:
            try:
                per100 = float(food.get(key) or 0.0)
            except (TypeError, ValueError):
                per100 = 0.0
            totals[key] += per100 * edible / 100.0
        resolved_grams += grams
        resolved.append(
            {
                "name": name,
                "normalized": norm,
                "food_code": str(food.get("food_code") or ""),
                "food_name": str(food.get("name") or ""),
                "grams": round(grams, 1),
                "edible_g": round(edible, 1),
                "stage": str(food.get("_stage") or ""),
            }
        )

    coverage = (resolved_grams / counted_grams) if counted_grams > 0 else 0.0
    return Estimate(
        nutrition={key: round(totals[key] / servings, 1) for key in NUTRIENTS},
        coverage=round(coverage, 4),
        servings=servings,
        resolved=resolved,
        unresolved=unresolved,
    )


# --- 成分表の取り込み（ADR-019 D1。DB を触る層） -----------------------------------------
#
# 列の位置は**固定でなく、ヘッダ名で探す**（八訂の Excel は版によって列が増える）。
# 見出しは複数行にまたがる（「エネルギー」の下に「kJ」「kcal」が並ぶ）ので、
# 列ごとに見出し行を**縦に連結**した文字列を作り、そこを語で探す。

#: 取り込む列の仕様: (契約の鍵, 必ず含む語, 含んでいたら外す語, 必須か)。
#: 「たんぱく質」は「アミノ酸組成によるたんぱく質」の列も語に当たるので `exclude` で外す
#: （どちらも残るときは**見出しが短い方**＝素の列を採る。`_pick_column` 参照）。
_COLUMN_SPECS: tuple[tuple[str, tuple[str, ...], tuple[str, ...], bool], ...] = (
    ("food_code", ("食品番号",), (), True),
    ("name", ("食品名",), (), True),
    ("food_group", ("食品群",), (), False),
    ("refuse_pct", ("廃棄率",), (), False),
    # 「エネルギー」は kJ 列との結合セルに載るので kcal 列の見出しには現れない。単位行の kcal だけで引く。
    ("kcal", ("kcal",), ("kj",), True),
    ("protein_g", ("たんぱく質",), ("アミノ酸",), True),
    ("fat_g", ("脂質",), ("脂肪酸", "トリアシル", "コレステロール"), True),
    ("carb_g", ("炭水化物",), ("利用可能", "単糖", "質量", "差引き", "食物繊維", "糖アルコール"), True),
    ("salt_g", ("食塩相当量",), (), True),
)

#: 見出しを探す行数の上限（八訂の本表は 10 行前後の飾りがある）。
_HEADER_SCAN_ROWS = 30


def _cell(row: Sequence[Any], i: int) -> str:
    if i < 0 or i >= len(row):
        return ""
    value = row[i]
    return "" if value is None else str(value).strip()


_HEADER_SPACES = re.compile(r"[\s\u3000]+")


def _header_cell(row: Sequence[Any], i: int) -> str:
    """見出し用の読み。八訂の見出しは「食　品　名」「廃　棄　率」のように全角空白で字を離すので、
    空白を全部落として語で引けるようにする（食品名そのものの読みには使わない）。"""
    return _HEADER_SPACES.sub("", _cell(row, i))


def _header_texts(rows: Sequence[Sequence[Any]], header_end: int) -> list[str]:
    """列ごとに、見出し行（0〜`header_end`）を縦に連結した文字列。"""
    width = max((len(r) for r in rows[: header_end + 1]), default=0)
    out: list[str] = []
    for i in range(width):
        parts = [_header_cell(rows[r], i) for r in range(header_end + 1)]
        out.append(" ".join(p for p in parts if p).lower())
    return out


def _pick_column(headers: Sequence[str], include: Sequence[str], exclude: Sequence[str]) -> int:
    """`include` を全部含み `exclude` を1つも含まない列のうち、見出しが**最も短い**もの。

    短い方を採るのは、八訂の見出しが「たんぱく質／アミノ酸組成によるたんぱく質」の
    ように**素の列の見出しが短く、派生列の見出しが長い**という規則性を持つため。
    """
    best = -1
    best_len = -1
    for i, text in enumerate(headers):
        if not all(word.lower() in text for word in include):
            continue
        if any(word.lower() in text for word in exclude):
            continue
        if best < 0 or len(text) < best_len:
            best, best_len = i, len(text)
    return best


def _looks_like_data_row(row: Sequence[Any]) -> bool:
    """食品番号らしい数字（3〜7桁）を持つ行か。見出しの終わりを見つけるのに使う。"""
    for i in range(len(row)):
        text = unicodedata.normalize("NFKC", _cell(row, i))
        if text.isdigit() and 3 <= len(text) <= 7:
            return True
    return False


def find_columns(rows: Sequence[Sequence[Any]]) -> tuple[int, dict[str, int]]:
    """見出しを探し、`(データの開始行, {契約の鍵: 列番号})` を返す。

    見出しは `食品名` が現れた行から始まり、**最初のデータ行の1つ前**で終わる
    ——八訂の本表は「エネルギー」の下に細目の行と単位の行（`kJ`／`kcal`）が来るので、
    何行あるかを決め打ちせず「食品番号らしい数字が出てきたら本体」と見る。
    こうすると列ごとの見出しが「エネルギー kcal」「たんぱく質 アミノ酸組成による g」の
    ように縦に連結でき、`kJ` の列と `kcal` の列を語で見分けられる。
    """
    name_row = -1
    for r, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        if any("食品名" in _header_cell(row, i) for i in range(len(row))):
            name_row = r
            break
    header_end = -1
    if name_row >= 0:
        header_end = name_row
        for r in range(name_row + 1, min(len(rows), _HEADER_SCAN_ROWS)):
            if _looks_like_data_row(rows[r]):
                break
            header_end = r
    if header_end < 0:
        raise ManorError(
            "成分表の見出し（「食品名」の列）が見つかりません。本表のシートを渡してください",
            code=2,
            key=ERR_FOOD_IMPORT_FAILED,
            params={"detail": "header"},
        )

    headers = _header_texts(rows, header_end)
    columns: dict[str, int] = {}
    missing: list[str] = []
    for key, include, exclude, required in _COLUMN_SPECS:
        i = _pick_column(headers, include, exclude)
        if i >= 0:
            columns[key] = i
        elif required:
            missing.append(include[0])
    if missing:
        raise ManorError(
            f"成分表に必要な列が見つかりません: {'、'.join(missing)}",
            code=2,
            key=ERR_FOOD_IMPORT_FAILED,
            params={"detail": ",".join(missing)},
        )
    return header_end + 1, columns


def _as_float(text: str) -> float | None:
    """成分表の数値。`Tr`（微量）・`-`（未測定）・`(0)`（推定値）を素直に読む。

    `Tr` と `-` は 0 として扱い、括弧つきの推定値は括弧を外して読む（成分表の約束）。
    """
    raw = unicodedata.normalize("NFKC", str(text or "")).strip()
    if not raw:
        return None
    if raw in ("Tr", "tr", "-", "−", "－"):
        return 0.0
    raw = raw.strip("()（）").replace(",", "")
    try:
        return float(raw)
    except ValueError:
        return None


def read_food_rows(path: Path) -> list[dict[str, Any]]:
    """Excel（`.xlsx`）または CSV から成分表の行を読む（ADR-019 D1）。

    `openpyxl` は**ここで遅延 import** する——取り込みのときだけ要るライブラリなので、
    `manor` の起動のたびに読み込まない（試験も CSV の偽データで済む）。
    """
    path = Path(path)
    if not path.is_file():
        raise ManorError(
            f"成分表のファイルが見つかりません: {path}",
            code=2,
            key=ERR_FOOD_IMPORT_FAILED,
            params={"detail": str(path)},
        )

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # 遅延 import（上の docstring 参照）
        except ImportError as exc:  # pragma: no cover - 依存が入っていない環境向け
            raise ManorError(
                "Excel を読むには openpyxl が必要です（`uv sync` を実行するか、CSV に変換して渡してください）",
                code=2,
                key=ERR_FOOD_IMPORT_FAILED,
                params={"detail": "openpyxl"},
            ) from exc
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            rows: list[Sequence[Any]] = [list(r) for r in sheet.iter_rows(values_only=True)]
        finally:
            book.close()
    else:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = [list(r) for r in csv.reader(f)]

    start, columns = find_columns(rows)
    out: list[dict[str, Any]] = []
    for row in rows[start:]:
        code = _cell(row, columns["food_code"])
        name = _cell(row, columns["name"])
        if not code or not name:
            continue
        if not unicodedata.normalize("NFKC", code).replace("-", "").isdigit():
            continue  # 単位の行・食品群の見出し行
        group = _cell(row, columns["food_group"]) if "food_group" in columns else ""
        out.append(
            {
                "food_code": unicodedata.normalize("NFKC", code),
                "food_group": group or unicodedata.normalize("NFKC", code)[:2],
                "name": unicodedata.normalize("NFKC", name),
                "kcal": _as_float(_cell(row, columns["kcal"])),
                "protein_g": _as_float(_cell(row, columns["protein_g"])),
                "fat_g": _as_float(_cell(row, columns["fat_g"])),
                "carb_g": _as_float(_cell(row, columns["carb_g"])),
                "salt_g": _as_float(_cell(row, columns["salt_g"])),
                "refuse_pct": (
                    _as_float(_cell(row, columns["refuse_pct"])) or 0.0
                    if "refuse_pct" in columns
                    else 0.0
                ),
            }
        )
    if not out:
        raise ManorError(
            "成分表に読める行がありませんでした（本表のシートを渡してください）",
            code=2,
            key=ERR_FOOD_IMPORT_FAILED,
            params={"detail": "empty"},
        )
    return out


def require_food_table(conn: sqlite3.Connection) -> None:
    """`chef_food` が無い home 向け（`cli._require_chef_recipe_table` と同じ流儀）。"""
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'chef_food'"
    ).fetchone()
    if row is None:
        raise ManorError(
            "食品成分表の表が未導入です。`manor init` を実行してください",
            code=2,
            key=ERR_FOOD_TABLE_MISSING,
        )


def import_food_table(
    conn: sqlite3.Connection, path: Path, *, source_version: str = DEFAULT_SOURCE_VERSION
) -> dict[str, object]:
    """成分表を `chef_food` へ取り込む。**冪等**（同じ `food_code` は上書き）。"""
    require_food_table(conn)
    rows = read_food_rows(Path(path))
    now = util.now()
    before = int(conn.execute("SELECT COUNT(*) AS n FROM chef_food").fetchone()["n"])
    for row in rows:
        conn.execute(
            "INSERT INTO chef_food"
            " (food_code, food_group, name, kcal, protein_g, fat_g, carb_g, salt_g,"
            "  refuse_pct, per, source_version, updated_at)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
            " ON CONFLICT (food_code) DO UPDATE SET"
            "  food_group = excluded.food_group, name = excluded.name, kcal = excluded.kcal,"
            "  protein_g = excluded.protein_g, fat_g = excluded.fat_g, carb_g = excluded.carb_g,"
            "  salt_g = excluded.salt_g, refuse_pct = excluded.refuse_pct,"
            "  per = excluded.per, source_version = excluded.source_version,"
            "  updated_at = excluded.updated_at",
            (
                row["food_code"], row["food_group"], row["name"], row["kcal"], row["protein_g"],
                row["fat_g"], row["carb_g"], row["salt_g"], row["refuse_pct"],
                PER_100G, source_version, now,
            ),
        )
    after = int(conn.execute("SELECT COUNT(*) AS n FROM chef_food").fetchone()["n"])
    return {
        "path": str(path),
        "rows": len(rows),
        "added": after - before,
        "updated": len(rows) - (after - before),
        "total": after,
        "source_version": source_version,
    }


# --- 成分表・名寄せの読み書き（DB を触る層） ----------------------------------------------


def food_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    """`chef_food` の全行（`build_index()` に渡す形）。数千行を超えない前提で素直に読む
    （`recipes.list_recipes` と同じ割り切り。八訂の本表は約 2,500 行）。"""
    require_food_table(conn)
    return [dict(r) for r in conn.execute("SELECT * FROM chef_food").fetchall()]


def alias_map(conn: sqlite3.Connection) -> dict[str, str]:
    """`chef_food_alias` の `alias → food_code`。"""
    require_food_table(conn)
    return {
        str(r["alias"]): str(r["food_code"])
        for r in conn.execute("SELECT alias, food_code FROM chef_food_alias").fetchall()
    }


def search_terms(q: str, tables: UnitTables | None = None) -> list[str]:
    """検索語 → LIKE に掛ける語の並び（ADR-019 §4）。**正規化の前と後の両方**を返す。

    成分表はかな書き（「たまねぎ」「ぶた」「こいくちしょうゆ」）、うちのレシピは漢字
    （「玉ねぎ」「豚」「醤油」）。`[food_normalize].synonyms` はレシピの書き方へ寄せる
    向きなので、検索では**逆向きにも当てる**——そうしないと「玉ねぎ」で成分表の
    「たまねぎ」が1件も出ない（2026-09-13 の実測）。
    """
    tables = tables or load_unit_tables()
    out: list[str] = []

    def push(value: str) -> None:
        text = str(value or "").strip()
        if text and text not in out:
            out.append(text)

    term = unicodedata.normalize("NFKC", str(q or "")).strip()
    push(term)
    push(shaping.normalize_food_name(term))
    norm = normalize_name(term, tables, drop=False)
    push(norm)
    for src, dst in tables.synonyms:
        key = shaping.normalize_food_name(dst)
        target = shaping.normalize_food_name(src)
        if key and target and key in norm:
            push(norm.replace(key, target))
    return out


def search_foods(conn: sqlite3.Connection, q: str, *, limit: int = 50) -> list[dict[str, Any]]:
    """食品名の**部分一致**で探す（画面の「食品を選ぶ」。ADR-019 D5）。

    引く語は `search_terms()` が広げる（正規化の前と後・同義語の逆向き）。どれか1つに
    当たれば拾い、**名前が短い順**に並べる——修飾の少ない素の食品が上に来る。
    """
    require_food_table(conn)
    terms = search_terms(q)
    if not terms:
        return []
    # LIKE の特殊文字（`%`・`_`）は主人が打った素の文字として扱う（ESCAPE で逃がす）。
    args: list[Any] = []
    for term in terms:
        args.append("%" + term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%")
    where = " OR ".join(["name LIKE ? ESCAPE '\\'"] * len(terms))
    args.append(max(1, int(limit)))
    rows = conn.execute(
        f"SELECT * FROM chef_food WHERE {where} ORDER BY LENGTH(name), food_code LIMIT ?",
        args,
    ).fetchall()
    return [dict(r) for r in rows]


def list_aliases(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    """登録済みの名寄せ（食品名を添える）。"""
    require_food_table(conn)
    rows = conn.execute(
        "SELECT a.alias, a.food_code, a.confidence, a.updated_at, f.name AS food_name"
        " FROM chef_food_alias a LEFT JOIN chef_food f ON f.food_code = a.food_code"
        " ORDER BY a.updated_at DESC, a.alias"
    ).fetchall()
    return [dict(r) for r in rows]


def set_alias(
    conn: sqlite3.Connection, alias: str, food_code: str, *, confidence: str = "manual"
) -> dict[str, Any]:
    """名寄せを1件入れる（冪等。同じ `alias` は上書き）。`alias` は正規化して持つ。"""
    require_food_table(conn)
    if confidence not in VALID_CONFIDENCE:
        raise ManorError(
            f"confidence は {'/'.join(VALID_CONFIDENCE)} のいずれかです: {confidence!r}",
            code=2,
            key=ERR_FOOD_NOT_FOUND,
            params={"detail": f"confidence={confidence!r}"},
        )
    norm = normalize_name(alias)
    if not norm:
        raise ManorError(
            "alias が空です", code=2, key=ERR_FOOD_NOT_FOUND, params={"detail": "alias"}
        )
    row = conn.execute(
        "SELECT food_code, name FROM chef_food WHERE food_code = ?", (str(food_code),)
    ).fetchone()
    if row is None:
        raise ManorError(
            f"食品が見つかりません: {food_code}",
            code=2,
            key=ERR_FOOD_NOT_FOUND,
            params={"detail": f"food_code={food_code!r}"},
        )
    now = util.now()
    conn.execute(
        "INSERT INTO chef_food_alias (alias, food_code, confidence, updated_at)"
        " VALUES (?, ?, ?, ?)"
        " ON CONFLICT (alias) DO UPDATE SET food_code = excluded.food_code,"
        "  confidence = excluded.confidence, updated_at = excluded.updated_at",
        (norm, str(food_code), confidence, now),
    )
    return {
        "alias": norm,
        "food_code": str(row["food_code"]),
        "food_name": str(row["name"]),
        "confidence": confidence,
        "updated_at": now,
    }


# --- 名寄せの種（ADR-019 §4） -------------------------------------------------------------


def load_alias_seed(
    path: Path | None = None, tables: UnitTables | None = None
) -> list[dict[str, Any]]:
    """`food_aliases_seed.toml` を読む（試験は `path` に合成データを渡せる）。

    戻り値は `[{"food_code", "name", "aliases": [正規化した材料名, …]}, …]`。
    **DB は触らない**——実在の検算（食品番号が `chef_food` にあるか）は `seed_aliases()`。
    """
    tables = tables or load_unit_tables()
    p = Path(path) if path is not None else _SEED_PATH
    try:
        with p.open("rb") as f:
            data = tomllib.load(f)
    except (OSError, tomllib.TOMLDecodeError) as exc:
        raise ManorError(
            f"名寄せの種を読めませんでした: {p}",
            code=2,
            key=ERR_FOOD_IMPORT_FAILED,
            params={"detail": str(p)},
        ) from exc
    out: list[dict[str, Any]] = []
    for raw in data.get("food") or []:
        if not isinstance(raw, Mapping):
            continue
        code = str(raw.get("code") or "").strip()
        if not code:
            continue
        aliases: list[str] = []
        for alias in raw.get("aliases") or []:
            norm = normalize_name(str(alias), tables)
            if norm and norm not in aliases:
                aliases.append(norm)
        if not aliases:
            continue
        out.append({"food_code": code, "name": str(raw.get("name") or ""), "aliases": aliases})
    return out


def seed_aliases(
    conn: sqlite3.Connection, *, path: Path | None = None, tables: UnitTables | None = None
) -> dict[str, object]:
    """名寄せの種を `chef_food_alias` へ `confidence='rule'` で入れる（冪等）。

    **`manual` は上書きしない**（ADR-019 §4）——画面や `manor chef food alias` で主人が
    決めた行は、種を入れ直しても動かない。同じ alias の `rule`／`llm` は更新する。

    `chef_food` に無い食品番号は**黙って飛ばし**（`missing` に数える）、例外にしない
    ——成分表を入れていない home や、版が違って番号が動いた行のために、種の1行で
    取り込み全体を失敗させたくない。
    """
    require_food_table(conn)
    tables = tables or load_unit_tables()
    entries = load_alias_seed(path, tables)
    known = {
        str(r["food_code"]) for r in conn.execute("SELECT food_code FROM chef_food").fetchall()
    }
    existing = {
        str(r["alias"]): str(r["confidence"])
        for r in conn.execute("SELECT alias, confidence FROM chef_food_alias").fetchall()
    }
    now = util.now()
    added = updated = kept = 0
    missing: list[str] = []
    seen: set[str] = set()
    for entry in entries:
        code = str(entry["food_code"])
        if code not in known:
            missing.append(code)
            continue
        for alias in entry["aliases"]:
            if alias in seen:
                continue  # 種の中の重複（試験が検算しているので、通れば起きない）
            seen.add(alias)
            if existing.get(alias) == "manual":
                kept += 1
                continue
            conn.execute(
                "INSERT INTO chef_food_alias (alias, food_code, confidence, updated_at)"
                " VALUES (?, ?, 'rule', ?)"
                " ON CONFLICT (alias) DO UPDATE SET food_code = excluded.food_code,"
                "  confidence = excluded.confidence, updated_at = excluded.updated_at",
                (alias, code, now),
            )
            if alias in existing:
                updated += 1
            else:
                added += 1
    return {
        "foods": len(entries),
        "aliases": len(seen),
        "added": added,
        "updated": updated,
        "kept_manual": kept,
        "missing": missing,
    }


def remove_alias(conn: sqlite3.Connection, alias: str) -> dict[str, object]:
    """名寄せを1件消す（無ければ何もしない。冪等）。"""
    require_food_table(conn)
    norm = normalize_name(alias)
    conn.execute("DELETE FROM chef_food_alias WHERE alias = ?", (norm,))
    return {"alias": norm, "removed": True}


# --- 再計算（ADR-019 D4「再計算の契機」） -------------------------------------------------


def _meta_row(conn: sqlite3.Connection, recipe_id: int) -> sqlite3.Row | None:
    return conn.execute(
        "SELECT nutrition_source FROM chef_recipe_meta WHERE recipe_id = ?", (recipe_id,)
    ).fetchone()


def _recipe_ids(conn: sqlite3.Connection, recipe_id: int | None) -> list[int]:
    if recipe_id is not None:
        return [int(recipe_id)]
    return [
        int(r["id"])
        for r in conn.execute(
            "SELECT id FROM chef_recipe WHERE archived_at IS NULL ORDER BY id"
        ).fetchall()
    ]


def estimate_for_recipe(
    conn: sqlite3.Connection,
    recipe_id: int,
    *,
    index: FoodIndex | None = None,
    aliases: Mapping[str, str] | None = None,
    tables: UnitTables | None = None,
) -> Estimate:
    """1 レシピを推定する（**書かない**）。画面・API の「見せるだけ」用。"""
    tables = tables or load_unit_tables()
    if index is None:
        index = build_index(food_rows(conn), tables)
    if aliases is None:
        aliases = alias_map(conn)
    recipe = recipes.get(conn, recipe_id)
    return estimate_nutrition(recipe, index, aliases, tables)


def rebuild(
    conn: sqlite3.Connection, *, recipe_id: int | None = None, quiet: bool = False
) -> dict[str, object]:
    """推定して `chef_recipe_meta` へ書く（ADR-019 D4）。

    `recipe_id` を省けば畳んでいないレシピ全部。**`site`／`manual` は上書きしない**
    ——人が確かめた数字を機械が塗り替えない（ADR-019 D4）。1つも解決できなかった
    レシピは**空欄のまま**にする（0 を書くと「栄養値がある」ことになってしまう）。

    `quiet=True` は表が無い home で静かに何もしない（登録・編集の経路に差し込むときに
    使う——成分表を入れていない主人のレシピ登録を、推定の都合で失敗させない）。
    """
    if quiet:
        row = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'chef_food'"
        ).fetchone()
        if row is None:
            return {"updated": 0, "skipped": 0, "unresolved": 0, "items": [], "available": False}
    tables = load_unit_tables()
    index = build_index(food_rows(conn), tables)
    aliases = alias_map(conn)
    minimum = coverage_min()

    updated = 0
    skipped = 0
    items: list[dict[str, object]] = []
    for rid in _recipe_ids(conn, recipe_id):
        meta = _meta_row(conn, rid)
        source = str(meta["nutrition_source"]) if meta is not None else ""
        if source in ("site", "manual"):
            skipped += 1
            items.append({"recipe_id": rid, "skipped": source})
            continue
        est = estimate_for_recipe(conn, rid, index=index, aliases=aliases, tables=tables)
        if not est.ok:
            skipped += 1
            items.append({"recipe_id": rid, "skipped": "no_match", "unresolved": len(est.unresolved)})
            continue
        recipes.set_meta(conn, rid, nutrition_source="estimated", **est.nutrition)
        conn.execute(
            "UPDATE chef_recipe_meta SET nutrition_coverage = ? WHERE recipe_id = ?",
            (round(est.coverage, 4), rid),
        )
        updated += 1
        items.append(
            {
                "recipe_id": rid,
                "coverage": round(est.coverage, 3),
                "partial": est.is_partial(minimum),
                "unresolved": len(est.unresolved),
                **est.nutrition,
            }
        )
    return {
        "updated": updated,
        "skipped": skipped,
        "unresolved": sum(int(i.get("unresolved") or 0) for i in items),
        "coverage_min": minimum,
        "items": items,
        "available": True,
    }


def refresh(conn: sqlite3.Connection, recipe_id: int) -> dict[str, object]:
    """1 レシピぶんの再計算（登録・材料の編集の直後に差し込む。静かに失敗する）。"""
    return rebuild(conn, recipe_id=recipe_id, quiet=True)


def nutrition_payload(conn: sqlite3.Connection, recipe_id: int) -> dict[str, object]:
    """`GET /api/v1/kitchen/recipes/{id}/nutrition` の応答（ADR-019 D5）。

    **保存されている値に `source`/`coverage`/`unresolved` を足すだけ**——XR（kitchen-xr）が
    読む5項目の形は変えない（ADR-019 D5「XR は今のまま」）。

    `coverage` は `chef_recipe_meta.nutrition_coverage`（推定のときだけ入っている）。
    `unresolved` は**その場で数え直す**——名寄せを直した直後に画面が新しい結果を
    見せられるようにする（保存しない。書くのは `rebuild()` だけ）。
    """
    recipe = recipes.get(conn, recipe_id)
    meta = dict(recipe["meta"])  # type: ignore[arg-type]
    source = str(meta.get("nutrition_source") or "")
    row = conn.execute(
        "SELECT nutrition_coverage FROM chef_recipe_meta WHERE recipe_id = ?", (recipe_id,)
    ).fetchone()
    stored_coverage = row["nutrition_coverage"] if row is not None else None

    available = (
        conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'chef_food'"
        ).fetchone()
        is not None
    )
    unresolved: list[dict[str, Any]] = []
    live_coverage: float | None = None
    if available and int(conn.execute("SELECT COUNT(*) AS n FROM chef_food").fetchone()["n"]) > 0:
        tables = load_unit_tables()
        est = estimate_nutrition(recipe, build_index(food_rows(conn), tables), alias_map(conn), tables)
        unresolved = est.unresolved
        live_coverage = est.coverage
    else:
        available = False

    coverage = stored_coverage if stored_coverage is not None else (
        live_coverage if source == "estimated" else None
    )
    minimum = coverage_min()
    return {
        "recipe_id": int(recipe["id"]),  # type: ignore[arg-type]
        "servings": recipe.get("servings"),
        **{key: meta.get(key) for key in NUTRIENTS},
        "source": source,
        "coverage": round(float(coverage), 3) if coverage is not None else None,
        "coverage_min": minimum,
        "partial": bool(source == "estimated" and coverage is not None and float(coverage) < minimum),
        "unresolved": unresolved,
        "food_table_available": available,
    }


def unresolved_summary(conn: sqlite3.Connection) -> dict[str, object]:
    """名寄せの画面（ADR-019 D5）の一覧。**未解決の材料名を束ねて**返す。

    同じ材料名は何本のレシピに出ても1行にまとめ（`count` と `recipes` を添える）、
    出る回数の多い順に並べる——手で名寄せするなら、効く順に見せたい。

    成分表をまだ入れていない home では空の一覧と `available: False` を返す
    ——画面は「成分表を取り込んでください」と案内できる（404 にしない。設定の画面が
    そこだけ壊れて見えるより、案内が出るほうがよい）。
    """
    require_food_table(conn)
    if int(conn.execute("SELECT COUNT(*) AS n FROM chef_food").fetchone()["n"]) == 0:
        return {"items": [], "total": 0, "available": False}
    tables = load_unit_tables()
    index = build_index(food_rows(conn), tables)
    aliases = alias_map(conn)
    groups: dict[str, dict[str, Any]] = {}
    for rid in _recipe_ids(conn, None):
        recipe = recipes.get(conn, rid)
        est = estimate_nutrition(recipe, index, aliases, tables)
        for item in est.unresolved:
            key = str(item.get("normalized") or item.get("name") or "")
            entry = groups.setdefault(
                key,
                {
                    "normalized": key,
                    "names": [],
                    "reason": item.get("reason"),
                    "count": 0,
                    "recipes": [],
                },
            )
            entry["count"] = int(entry["count"]) + 1
            if item["name"] not in entry["names"]:
                entry["names"].append(item["name"])
            if len(entry["recipes"]) < 5:
                entry["recipes"].append({"recipe_id": rid, "title": str(recipe.get("title") or "")})
    items = sorted(groups.values(), key=lambda e: (-int(e["count"]), str(e["normalized"])))
    return {"items": items, "total": len(items)}
